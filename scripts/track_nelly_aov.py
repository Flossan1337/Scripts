#!/usr/bin/env python3

import os
import csv
import re
import time
import statistics
from datetime import date
from bs4 import BeautifulSoup

from selenium import webdriver
from selenium.webdriver.chrome.service import Service
from selenium.webdriver.chrome.options import Options
from webdriver_manager.chrome import ChromeDriverManager

# NEW: Excel
from openpyxl import Workbook, load_workbook
from pathlib import Path

# ──────────────────────────────────────────────────────────────────────────────
URLS = [
    "https://nelly.com/se/topplistan/?page=1",
    "https://nelly.com/se/topplistan/?page=2",
    "https://nelly.com/se/topplistan/?page=3",
    "https://nelly.com/se/topplistan/?page=4",
    "https://nelly.com/se/topplistan/?page=5",
    "https://nelly.com/se/topplistan/?page=6",
    "https://nelly.com/se/topplistan/?page=7",
    "https://nelly.com/se/topplistan/?page=8",
    "https://nelly.com/se/topplistan/?page=9",
    "https://nelly.com/se/topplistan/?page=10",
]

# If a product is discounted, Nelly renders the selling price inside <ins> and the
# original price inside <del>. Non-discounted items show the price in a span like
# <span class="text-sm text-darkGrey">299&nbsp;kr</span>.
DISCOUNT_PRICE_SELECTOR = "ins"
REGULAR_PRICE_SELECTOR = "span.text-sm.text-darkGrey"

# OLD: CSV_FILENAME = "nelly_aov.csv"
# NEW: write to ../data/nelly_aov.xlsx (repo-root/data)
SCRIPT_DIR = Path(__file__).resolve().parent
XLSX_PATH = (SCRIPT_DIR / ".." / "data" / "nelly_aov.xlsx").resolve()
# ──────────────────────────────────────────────────────────────────────────────

def create_driver():
    opts = Options()
    opts.add_argument("--headless")
    opts.add_argument("--window-size=1920,1080")
    opts.add_argument("--disable-gpu")
    opts.add_argument("--no-sandbox")
    opts.add_argument(
        "user-agent=Mozilla/5.0 (Windows NT 10.0; Win64; x64) "
        "AppleWebKit/537.36 (KHTML, like Gecko) "
        "Chrome/115.0.0.0 Safari/537.36"
    )
    service = Service(ChromeDriverManager().install())
    return webdriver.Chrome(service=service, options=opts)

def _to_number(text):
    # Keep digits, comma, dot; normalize comma to dot; strip NBSPs.
    num = re.sub(r"[^\d,\.]", "", text.replace("\xa0", " "))
    num = num.replace(",", ".")
    try:
        return float(num)
    except ValueError:
        return None

def fetch_prices(driver, url):
    driver.get(url)
    time.sleep(3)
    # Trigger lazy loading
    driver.execute_script("window.scrollTo(0, document.body.scrollHeight);")
    time.sleep(2)

    soup = BeautifulSoup(driver.page_source, "html.parser")

    prices = []

    # 1) Collect all discounted selling prices (<ins>)
    for tag in soup.select(DISCOUNT_PRICE_SELECTOR):
        val = _to_number(tag.get_text(strip=True))
        if val is not None:
            prices.append(val)

    # 2) Collect all regular prices only for items that don't show a discount.
    # Heuristic: regular price spans exist broadly; to avoid double counting
    # discounted items we only add spans when the *nearest* container does not
    # also contain an <ins>. This keeps things robust even if the DOM shifts a bit.
    for span in soup.select(REGULAR_PRICE_SELECTOR):
        container = span
        # Climb a few levels up looking for an <ins> sibling in the same card.
        has_discount = False
        for _ in range(4):
            if container is None or container.parent is None:
                break
            container = container.parent
            if container.select_one(DISCOUNT_PRICE_SELECTOR):
                has_discount = True
                break
        if has_discount:
            continue

        val = _to_number(span.get_text(strip=True))
        if val is not None:
            prices.append(val)

    return prices

# ──────────────────────────────────────────────────────────────────────────────
# Excel handling (replaces CSV)

def ensure_header_xlsx():
    XLSX_PATH.parent.mkdir(parents=True, exist_ok=True)
    if XLSX_PATH.exists():
        return
    wb = Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    ws.append(["date", "median_price", "average_price"])
    wb.save(XLSX_PATH)

def append_to_xlsx(median_price, average_price):
    today = date.today().isoformat()
    if not XLSX_PATH.exists():
        ensure_header_xlsx()
    wb = load_workbook(XLSX_PATH)
    ws = wb.active
    ws.append([today, round(median_price, 2), round(average_price, 2)])
    wb.save(XLSX_PATH)
    print(f"✓ Appended {today}: median={median_price:.2f}, avg={average_price:.2f} → {XLSX_PATH}")

# ──────────────────────────────────────────────────────────────────────────────

def main():
    ensure_header_xlsx()
    driver = create_driver()
    all_prices = []

    for url in URLS:
        ps = fetch_prices(driver, url)
        print(f"  • Nelly Topplistan – found {len(ps)} prices on {url}")
        all_prices.extend(ps)

    driver.quit()

    if not all_prices:
        print("❌ No prices found; check selectors or if the site changed its HTML.")
        return

    med = statistics.median(all_prices)
    avg = statistics.mean(all_prices)
    append_to_xlsx(med, avg)

if __name__ == "__main__":
    main()
